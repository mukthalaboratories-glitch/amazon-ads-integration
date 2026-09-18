import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from collections import defaultdict

BASE = os.path.dirname(os.path.abspath(__file__))
PROJECT = os.path.dirname(BASE)

with open(os.path.join(PROJECT, 'amazon_july_raw.json'), 'r', encoding='utf-8') as f:
    amazon = json.load(f)

wb = openpyxl.load_workbook(os.path.join(PROJECT, 'Firstcry sales', 'dashboardsale_2026_07_01_to_2026_07_31 (1).xlsx'), data_only=True)
ws_fc = wb['Sales Data']
firstcry = []
for r in ws_fc.iter_rows(min_row=2, values_only=True):
    poid, dt, pid, brand, btype, qty, mrp, sales, subcat, cat, stock, sku = r
    firstcry.append({'po': poid, 'day': dt[:10], 'sku': sku, 'product_id': pid,
                     'qty': qty, 'sales': sales})

flipkart = [
    {'po': 'OD338234475046705100', 'day': '2026-07-31', 'product': 'Catche Insta Ayurvedic Mosquito Repellent (Pack of 8)',
     'sku': 'Catche must-quit-o Insta Repellent (Pack of 8)', 'qty': 1, 'sales': 425},
]

# ---- Aggregate by day ----
day_data = defaultdict(lambda: {'amazon_sales': 0, 'amazon_units': 0, 'amazon_orders': 0,
                                'fc_sales': 0, 'fc_units': 0, 'fc_orders': 0,
                                'fk_sales': 0, 'fk_units': 0, 'fk_orders': 0})

for r in amazon:
    d = day_data[r['day']]
    d['amazon_sales'] += r['sales']
    d['amazon_units'] += r['units']
    d['amazon_orders'] += 1

for r in firstcry:
    d = day_data[r['day']]
    d['fc_sales'] += r['sales']
    d['fc_units'] += r['qty']
    d['fc_orders'] += 1

for r in flipkart:
    d = day_data[r['day']]
    d['fk_sales'] += r['sales']
    d['fk_units'] += r['qty']
    d['fk_orders'] += 1

# ---- Aggregates ----
def total(key):
    return sum(d[key] for d in day_data.values())

totals = {
    'amazon_sales': total('amazon_sales'), 'amazon_units': total('amazon_units'), 'amazon_orders': total('amazon_orders'),
    'fc_sales': total('fc_sales'), 'fc_units': total('fc_units'), 'fc_orders': total('fc_orders'),
    'fk_sales': total('fk_sales'), 'fk_units': total('fk_units'), 'fk_orders': total('fk_orders'),
}
totals['total_sales'] = totals['amazon_sales'] + totals['fc_sales'] + totals['fk_sales']
totals['total_units'] = totals['amazon_units'] + totals['fc_units'] + totals['fk_units']
totals['total_orders'] = totals['amazon_orders'] + totals['fc_orders'] + totals['fk_orders']

# ---- Product level ----
amazon_products = defaultdict(lambda: {'units': 0, 'sales': 0, 'orders': set()})
for r in amazon:
    key = (r['asin'], r['sku'], r['title'])
    p = amazon_products[key]
    p['units'] += r['units']
    p['sales'] += r['sales']
    if r['orderId']:
        p['orders'].add(r['orderId'])

fc_products = defaultdict(lambda: {'units': 0, 'sales': 0, 'orders': set()})
for r in firstcry:
    key = ('', r['sku'], f"Catche Must-quit-o (PID {r['product_id']})")
    p = fc_products[key]
    p['units'] += r['qty']
    p['sales'] += r['sales']
    if r['po']:
        p['orders'].add(r['po'])

fk_products = defaultdict(lambda: {'units': 0, 'sales': 0, 'orders': set()})
for r in flipkart:
    key = ('', r['sku'], r['product'])
    p = fk_products[key]
    p['units'] += r['qty']
    p['sales'] += r['sales']
    if r['po']:
        p['orders'].add(r['po'])

# ---- Excel build ----
out_path = os.path.join(PROJECT, 'docs', 'channel-sales-july-2026.xlsx')
os.makedirs(os.path.dirname(out_path), exist_ok=True)

wb_out = openpyxl.Workbook()

header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_font = Font(color='FFFFFF', bold=True)
center = Alignment(horizontal='center')

def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center

# Sheet 1: Summary
ws = wb_out.active
ws.title = 'Summary'
ws.append(['Channel', 'Orders', 'Units', 'Sales (INR)', 'Avg Order Value (INR)'])
channels = [
    ('Amazon.in', totals['amazon_orders'], totals['amazon_units'], totals['amazon_sales']),
    ('FirstCry', totals['fc_orders'], totals['fc_units'], totals['fc_sales']),
    ('Flipkart', totals['fk_orders'], totals['fk_units'], totals['fk_sales']),
]
for name, o, u, s in channels:
    ws.append([name, o, u, round(s, 2), round(s / o, 2) if o else 0])
ws.append(['TOTAL', totals['total_orders'], totals['total_units'], round(totals['total_sales'], 2),
           round(totals['total_sales'] / totals['total_orders'], 2) if totals['total_orders'] else 0])
style_header(ws, 1, 5)
for col, w in zip('ABCDE', [18, 10, 10, 14, 20]):
    ws.column_dimensions[col].width = w
ws.cell(row=len(channels) + 2, column=1).font = Font(bold=True)

# Sheet 2: Daily breakdown
ws2 = wb_out.create_sheet('Daily')
ws2.append(['Date', 'Amazon Orders', 'Amazon Units', 'Amazon Sales', 'FirstCry Orders', 'FirstCry Units', 'FirstCry Sales',
            'Flipkart Orders', 'Flipkart Units', 'Flipkart Sales', 'Total Orders', 'Total Units', 'Total Sales'])
for day in sorted(day_data):
    d = day_data[day]
    ws2.append([
        day,
        d['amazon_orders'], d['amazon_units'], round(d['amazon_sales'], 2),
        d['fc_orders'], d['fc_units'], round(d['fc_sales'], 2),
        d['fk_orders'], d['fk_units'], round(d['fk_sales'], 2),
        d['amazon_orders'] + d['fc_orders'] + d['fk_orders'],
        d['amazon_units'] + d['fc_units'] + d['fk_units'],
        round(d['amazon_sales'] + d['fc_sales'] + d['fk_sales'], 2),
    ])
style_header(ws2, 1, 13)
for col, w in zip('ABCDEFGHIJKLM', [12, 13, 12, 13, 14, 13, 13, 14, 13, 13, 13, 12, 12]):
    ws2.column_dimensions[col].width = w

# Sheet 3: Amazon products
ws3 = wb_out.create_sheet('Amazon Products')
ws3.append(['ASIN', 'SKU', 'Product', 'Units', 'Sales (INR)', 'Orders'])
for (asin, sku, title), p in sorted(amazon_products.items(), key=lambda x: -x[1]['sales']):
    ws3.append([asin, sku, title, p['units'], round(p['sales'], 2), len(p['orders'])])
style_header(ws3, 1, 6)
for col, w in zip('ABCDEF', [14, 32, 60, 10, 14, 10]):
    ws3.column_dimensions[col].width = w

# Sheet 4: FirstCry products
ws4 = wb_out.create_sheet('FirstCry Products')
ws4.append(['SKU', 'Product', 'Units', 'Sales (INR)', 'Orders'])
for (asin, sku, title), p in sorted(fc_products.items(), key=lambda x: -x[1]['sales']):
    ws4.append([sku, title, p['units'], round(p['sales'], 2), len(p['orders'])])
style_header(ws4, 1, 5)
for col, w in zip('ABCDE', [32, 45, 10, 14, 10]):
    ws4.column_dimensions[col].width = w

# Sheet 5: Flipkart products
ws5 = wb_out.create_sheet('Flipkart Products')
ws5.append(['SKU', 'Product', 'Units', 'Sales (INR)', 'Orders'])
for (asin, sku, title), p in sorted(fk_products.items(), key=lambda x: -x[1]['sales']):
    ws5.append([sku, title, p['units'], round(p['sales'], 2), len(p['orders'])])
style_header(ws5, 1, 5)
for col, w in zip('ABCDE', [40, 50, 10, 14, 10]):
    ws5.column_dimensions[col].width = w

wb_out.save(out_path)
print(f'Excel saved: {out_path}')
print()
print('=== TOTALS (July 2026) ===')
for name, o, u, s in channels:
    print(f'{name}: {o} orders | {u} units | Rs{s:,.2f}')
print(f'TOTAL: {totals["total_orders"]} orders | {totals["total_units"]} units | Rs{totals["total_sales"]:,.2f}')
print()
print('=== JULY 31 ===')
if '2026-07-31' in day_data:
    d = day_data['2026-07-31']
    print(f'Amazon:   {d["amazon_orders"]} orders | {d["amazon_units"]} units | Rs{d["amazon_sales"]:,.2f}')
    print(f'FirstCry: {d["fc_orders"]} orders | {d["fc_units"]} units | Rs{d["fc_sales"]:,.2f}')
    print(f'Flipkart: {d["fk_orders"]} orders | {d["fk_units"]} units | Rs{d["fk_sales"]:,.2f}')
    t = d['amazon_orders'] + d['fc_orders'] + d['fk_orders']
    u = d['amazon_units'] + d['fc_units'] + d['fk_units']
    s = d['amazon_sales'] + d['fc_sales'] + d['fk_sales']
    print(f'TOTAL:    {t} orders | {u} units | Rs{s:,.2f}')
