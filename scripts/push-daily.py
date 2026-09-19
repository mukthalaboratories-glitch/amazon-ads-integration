import json
import os
import sys
import requests
from collections import defaultdict
from openpyxl import load_workbook

BASE = os.path.dirname(os.path.abspath(__file__))
PROJECT = os.path.dirname(BASE)

WEB_APP_URL = 'https://script.google.com/macros/s/AKfycbz07KglX7GGgM6JCbaYxk78S7JTV4gcSNRVR8Da5PM-e5G4w-QyscZqOvwkdx34nuZe/exec'
PUSH_KEY = 'catche-daily-2026'

FC_CATALOG = {
    '16071685': ('8908006802180', 368, 'Catche Ayurvedic Mosquito Repellent Refill (Pack of 4)'),
    '16071684': ('8908006802159', 225, 'Catche Ayurvedic Mosquito Repellent Combo (2 Refills + 1 Machine)'),
    '16071689': ('8908006802319', 360, 'Catche Herbal Mosquito Incense Sticks (120 sticks)'),
    '16071686': ('8908006802005', 450, 'Catche Ayurvedic Mosquito Trapellent Combo (Device + Refill)'),
    '16071687': ('8908006802180', 300, 'Catche Ayurvedic Mosquito Trapellent Refill (Pack of 4)'),
    '16071688': ('8908006802135', 270, 'Catche Ayurvedic Mosquito Lotion (60ml) (Pack of 3)'),
    '16071690': ('8908006802043', 295, 'Gadiva Ayurvedic Hairoil (100ml) with Neem Comb'),
}

FLIPKART_BY_DAY = {
    '2026-07-31': [
        {'po': 'OD338234475046705100',
         'product': 'Catche Insta Ayurvedic Mosquito Repellent (Pack of 8)',
         'sku': 'Catche must-quit-o Insta Repellent (Pack of 8)', 'qty': 1, 'sales': 425},
    ],
    '2026-08-01': [
        {'po': 'OD438243036104837100',
         'product': 'Catche Insta Ayurvedic Mosquito Repellent (Pack of 8)',
         'sku': 'Catche must-quit-o Insta Repellent (Pack of 8)', 'qty': 1, 'sales': 421},
        {'po': 'OD338245037653222100',
         'product': 'Catche Insta Ayurvedic Mosquito Repellent (Pack of 8)',
         'sku': 'Catche must-quit-o Insta Repellent (Pack of 8)', 'qty': 1, 'sales': 421},
        {'po': 'OD338245851671914100',
         'product': 'Catche Ayurvedic Mosquito Trapellent Combo (Device + Refill)',
         'sku': 'CT_TRAP_COMBO', 'qty': 1, 'sales': 415},
        {'po': 'OD338245441398294100',
         'product': 'Catche Must-Quit-O Insta Combo Mosquito Vaporizer (1 Machine + 2 Refills)',
         'sku': 'CT_INSTA_COM_1M+2R', 'qty': 1, 'sales': 184},
    ],
    '2026-08-02': [
        {'po': 'OD338248251862289100',
         'product': 'Catche Herbal Mosquito Incense Sticks (120 sticks)',
         'sku': '8908006802319', 'qty': 1, 'sales': 179},
        {'po': 'OD438251258242377100',
         'product': 'Catche Insta Ayurvedic Mosquito Repellent (Pack of 6)',
         'sku': 'Catche Insta Ayurvedic Mosquito Repellent-6pack', 'qty': 1, 'sales': 345},
        {'po': 'OD438253557867941100',
         'product': 'Catche Must-Quit-O Insta Combo Mosquito Vaporizer (1 Machine + 2 Refills)',
         'sku': 'CT_INSTA_COM_1M+2R', 'qty': 1, 'sales': 184},
    ],
    '2026-08-03': [
        {'po': 'OD438258052393795400',
         'product': 'Catche Herbal Mosquito Incense Sticks (120 sticks)',
         'sku': '8908006802319', 'qty': 3, 'sales': 599},
        {'po': 'OD338262381891458100',
         'product': 'Catche Herbal Mosquito Incense Sticks (120 sticks)',
         'sku': '8908006802319', 'qty': 1, 'sales': 179},
        {'po': 'OD438257341037844100',
         'product': 'Catche Must-Quit-O Insta Combo Mosquito Vaporizer (1 Machine + 2 Refills)',
         'sku': 'CT_INSTA_COM_1M+2R', 'qty': 1, 'sales': 176},
        {'po': 'OD438262082069101100',
         'product': 'Catche Must-Quit-O Insta Combo Mosquito Vaporizer (1 Machine + 2 Refills)',
         'sku': 'CT_INSTA_COM_1M+2R', 'qty': 1, 'sales': 192},
        {'po': 'OD438262544710897100',
         'product': 'Catche Insta Ayurvedic Mosquito Repellent (Pack of 6)',
         'sku': 'Catche Insta Ayurvedic Mosquito Repellent-6pack', 'qty': 1, 'sales': 345},
        {'po': 'OD338261909861569100',
         'product': 'Catche Insta Ayurvedic Mosquito Repellent (Pack of 8)',
         'sku': 'Catche must-quit-o Insta Repellent (Pack of 8)', 'qty': 1, 'sales': 411},
        {'po': 'OD438262851813781100',
         'product': 'Catche Must-Quit-O Insta Combo Mosquito Vaporizer (1 Machine + 2 Refills)',
         'sku': 'CT_INSTA_COM_1M+2R', 'qty': 1, 'sales': 184},
    ],
    '2026-08-05': [
        {'po': 'OD338279662655574100',
         'product': 'Catche Insta Ayurvedic Mosquito Repellent (Pack of 8)',
         'sku': 'Catche must-quit-o Insta Repellent (Pack of 8)', 'qty': 1, 'sales': 421},
        {'po': 'OD438279579357578100',
         'product': 'Catche Insta Ayurvedic Mosquito Repellent (Pack of 8)',
         'sku': 'Catche must-quit-o Insta Repellent (Pack of 8)', 'qty': 1, 'sales': 421},
    ],
}


import glob as _glob
# auto-detect latest FirstCry/Flipkart sales (dedupe by POID/order_item_id handles overlaps)
FC_FILES = sorted([os.path.basename(p) for p in _glob.glob(os.path.join(PROJECT, 'Firstcry sales', 'dashboardsale*.xlsx'))])
if not FC_FILES:
    FC_FILES = [
        'dashboardsale_2026_06_01_to_2026_06_30 (1).xlsx',
        'dashboardsale_2026_07_01_to_2026_07_31 (1).xlsx',
        'dashboardsale_2026_08_01_to_2026_08_31 (7).xlsx',
        'dashboardsale_2026_09_01_to_2026_09_30.xlsx',
    ]

FC_MAX_DAY = '2099-12-31'  # no cap — files themselves define range, dedupe by POID handles overlaps

FK_SALES_FILES = sorted([os.path.join('flipkart sales', os.path.basename(p)) for p in _glob.glob(os.path.join(PROJECT, 'flipkart sales', '*.xlsx'))])
if not FK_SALES_FILES:
    FK_SALES_FILES = [r'flipkart sales\1fe490a7-53b7-4ad4-8bc8-7edb75b47703_1788328552000.xlsx']

FK_PRICE_BY_SKU = {
    'Catche must-quit-o Insta Repellent (Pack of 8)': 420,
    'CT_INSTA_COM_1M+2R': 184,
    'CT_TRAP_COMBO': 415,
    '8908006802319': 186,
    'Catche Insta Ayurvedic Mosquito Repellent-6pack': 345,
    'CT_INSTA_COM_1M+1R': 250,
    'CT_INSTA_REF_P4': 337,
    'CT_INSTA_REF_P2': 119,
    'GD_HAIR_OIL': 249,
    'MT_SCRUBBER_SHEET': 150,
    'Catche mosquito repellent lotion - 60ml(Pack of 3)': 270,
    'Catche mosquito insta repellent combo': 184,
    'Catche must-quit-o Refill (pack of 8)': 420,
    'MT_HW_175ML': 150,
    'Muktha Ayurvedic Handwash': 150,
}

FK_CANONICAL_NAME = {
    'Catche must-quit-o Insta Repellent (Pack of 8)': 'Catche Insta Ayurvedic Mosquito Repellent (Pack of 8)',
    'CT_TRAP_COMBO': 'Catche Ayurvedic Mosquito Trapellent Combo (Device + Refill)',
    'CT_INSTA_COM_1M+2R': 'Catche Must-Quit-O Insta Combo Mosquito Vaporizer (1 Machine + 2 Refills)',
    '8908006802319': 'Catche Herbal Mosquito Incense Sticks (120 sticks)',
    'Catche Insta Ayurvedic Mosquito Repellent-6pack': 'Catche Insta Ayurvedic Mosquito Repellent (Pack of 6)',
}


def load_firstcry():
    firstcry = []
    seen_po = set()
    for fname in FC_FILES:
        fpath = os.path.join(PROJECT, 'Firstcry sales', fname)
        if not os.path.exists(fpath):
            continue
        wb = load_workbook(fpath, data_only=True)
        ws_fc = wb['Sales Data']
        for r in ws_fc.iter_rows(min_row=2, values_only=True):
            poid, dt, pid, brand, btype, qty, mrp, sales, subcat, cat, stock, sku = r
            po_key = f'{poid}|{pid}'
            if po_key in seen_po:
                continue
            seen_po.add(po_key)
            firstcry.append({'po': poid, 'day': dt[:10], 'product_id': str(pid),
                             'qty': qty, 'sales': sales})
    return firstcry


def load_flipkart_sales():
    fk = []
    seen_item = set()
    known_pos = {r['po'] for rows in FLIPKART_BY_DAY.values() for r in rows}
    for rel in FK_SALES_FILES:
        fpath = os.path.join(PROJECT, rel)
        if not os.path.exists(fpath):
            continue
        wb = load_workbook(fpath, data_only=True)
        ws_fk = wb['Orders']
        hdr = [c.value for c in ws_fk[1]]
        idx = {h: i for i, h in enumerate(hdr)}
        for r in ws_fk.iter_rows(min_row=2, values_only=True):
            item_id = str(r[idx['order_item_id']])
            if item_id in seen_item:
                continue
            seen_item.add(item_id)
            order_id = str(r[idx['order_id']] or '').strip()
            if order_id in known_pos:
                continue
            sku = str(r[idx['sku']] or '').strip('"').replace('SKU:', '').strip()
            title = str(r[idx['product_title']] or '').strip().strip('"').strip()
            qty = int(r[idx['quantity']] or 0)
            day = str(r[idx['order_date']])[:10]
            price = FK_PRICE_BY_SKU.get(sku)
            if price is None:
                print(f'WARNING: no price for SKU {sku!r} ({title[:40]}), using 0')
                price = 0
            name = FK_CANONICAL_NAME.get(sku, title)
            fk.append({'po': order_id, 'day': day, 'sku': sku, 'product': name,
                       'qty': qty, 'sales': round(price * qty, 2)})
    return fk


def build_rows(day, channels=None):
    amazon = []
    for fname in ('amazon_june_raw.json', 'amazon_july_raw.json', 'amazon_aug_raw.json', 'amazon_sep_raw.json'):
        fpath = os.path.join(PROJECT, fname)
        if os.path.exists(fpath):
            with open(fpath, 'r', encoding='utf-8') as f:
                amazon += json.load(f)

    firstcry = load_firstcry()

    amazon_products = defaultdict(lambda: {'units': 0, 'sales': 0, 'orders': set(), 'state': '', 'city': ''})
    for r in amazon:
        if r['day'] != day:
            continue
        if not str(r.get('title', '')).strip() or r['title'].strip() == '-':
            continue
        key = (r['asin'], r['sku'], r['title'].strip(), (r.get('shipState') or '').upper(), (r.get('shipCity') or '').upper())
        p = amazon_products[key]
        p['units'] += r['units']
        p['sales'] += r['sales']
        p['state'] = p['state'] or r.get('shipState', '')
        p['city'] = p['city'] or r.get('shipCity', '')
        if r['orderId']:
            p['orders'].add(r['orderId'])

    fc_products = defaultdict(lambda: {'units': 0, 'sales': 0, 'orders': set(), 'name': ''})
    for r in firstcry:
        if r['day'] != day:
            continue
        if day > FC_MAX_DAY:
            continue
        pid = r['product_id']
        cat = FC_CATALOG.get(pid, ('', 0, f'Catche Must-quit-o (PID {pid})'))
        p = fc_products[pid]
        p['units'] += r['qty']
        p['sales'] += r['sales']
        p['name'] = cat[2]
        if r['po']:
            p['orders'].add(r['po'])

    fk_products = defaultdict(lambda: {'units': 0, 'sales': 0, 'orders': set()})
    fk_rows = list(FLIPKART_BY_DAY.get(day, [])) + [r for r in load_flipkart_sales() if r['day'] == day]
    for r in fk_rows:
        key = (r['sku'], r['product'])
        p = fk_products[key]
        p['units'] += r['qty']
        p['sales'] += r['sales']
        if r.get('po'):
            p['orders'].add(r['po'])
        else:
            p['orders'].add(day)

    rows = []
    if not channels or 'amazon' in channels:
        for (asin, sku, title, _state, _city), p in amazon_products.items():
            rows.append({'date': day, 'channel': 'Amazon', 'product': title, 'asin': asin, 'sku': sku,
                         'units': p['units'], 'sales': round(p['sales'], 2), 'orders': len(p['orders']),
                         'status': 'Shipped', 'shipState': p['state'], 'shipCity': p['city']})
    if not channels or 'firstcry' in channels:
        for pid, p in fc_products.items():
            rows.append({'date': day, 'channel': 'FirstCry', 'product': p['name'], 'asin': '',
                         'sku': f'PID {pid}', 'units': p['units'], 'sales': round(p['sales'], 2),
                         'orders': len(p['orders']), 'status': '', 'shipState': '', 'shipCity': ''})
    if not channels or 'flipkart' in channels:
        for (sku, prod), p in fk_products.items():
            rows.append({'date': day, 'channel': 'Flipkart', 'product': prod, 'asin': '', 'sku': sku,
                         'units': p['units'], 'sales': round(p['sales'], 2), 'orders': len(p['orders']),
                         'status': '', 'shipState': '', 'shipCity': ''})
    return rows


def get_json(url, params, attempts=4, timeout=120):
    import time
    last = None
    for i in range(attempts):
        try:
            resp = requests.get(url, params=params, timeout=timeout)
            if resp.status_code == 200 and resp.text.strip():
                return resp.json()
            last = 'empty response (HTTP {})'.format(resp.status_code)
        except Exception as e:
            last = str(e)
        time.sleep(10)
    raise RuntimeError('GET failed after {} attempts: {}'.format(attempts, last))


def post_json(url, payload, attempts=12, timeout=120):
    import time
    last = None
    for i in range(attempts):
        try:
            resp = requests.post(url, data=json.dumps(payload).encode('utf-8'),
                                 headers={'Content-Type': 'application/json'}, timeout=timeout)
            if resp.status_code == 200 and resp.text.strip():
                return resp.json()
            last = 'empty response (HTTP {})'.format(resp.status_code)
        except Exception as e:
            last = str(e)
        time.sleep(20)
    raise RuntimeError('POST failed after {} attempts: {}'.format(attempts, last))


def main():
    day = sys.argv[1] if len(sys.argv) > 1 else '2026-07-31'
    channels = sys.argv[2].split(',') if len(sys.argv) > 2 else None
    rows = build_rows(day, channels)
    total_units = sum(r['units'] for r in rows)
    total_sales = sum(r['sales'] for r in rows)
    print(f'Built {len(rows)} rows for {day} '
          f'({total_units} units, Rs {total_sales:,.2f})')

    existing = None
    try:
        existing = get_json(WEB_APP_URL, {'key': PUSH_KEY, 'action': 'all'}, attempts=6, timeout=90)
    except Exception:
        print('Warning: GET existing failed, relying on server-side dedupe')
    existing_keys = set()
    if existing:
        for r in existing.get('rows', []):
            existing_keys.add(row_key(r))

    new_rows = []
    for r in rows:
        if row_key(r) in existing_keys:
            continue
        existing_keys.add(row_key(r))
        new_rows.append(r)

    if not new_rows:
        print('Nothing to push - all rows already in sheet (dedupe OK)')
        return

    print(f'Pushing {len(new_rows)} new rows '
          f'({sum(r["units"] for r in new_rows)} units, '
          f'Rs {sum(r["sales"] for r in new_rows):,.2f})')

    result = post_json(WEB_APP_URL, {'key': PUSH_KEY, 'rows': new_rows})
    print(result)


def row_key(r):
    return '|'.join([
        str(r.get('date', '')),
        str(r.get('channel', '')),
        str(r.get('asin') or r.get('product') or ''),
        str(r.get('product', '')),
        str(r.get('sku', '')),
        str(r.get('status', '')),
        str(r.get('shipState', '')),
        str(r.get('shipCity', '')),
    ]).lower()


if __name__ == '__main__':
    main()



